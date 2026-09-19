from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[1]
LECTURES = ROOT / "Лекції"
OUT = ROOT / ".analysis_tmp" / "lecture_corpus.json"


def clean(value: object) -> str:
    return " ".join(str(value).replace("\x0b", "\n").split())


def extract_chart(chart) -> dict:
    result: dict = {"title": "", "series": []}
    if chart.has_title and chart.chart_title.has_text_frame:
        result["title"] = clean(chart.chart_title.text_frame.text)
    for series in chart.series:
        entry = {"name": clean(series.name), "values": []}
        try:
            entry["values"] = [v for v in series.values]
        except Exception:
            pass
        result["series"].append(entry)
    try:
        plots = []
        for plot in chart.plots:
            cats = []
            try:
                cats = [clean(c.label) for c in plot.categories]
            except Exception:
                pass
            plots.append({"categories": cats})
        result["plots"] = plots
    except Exception:
        pass
    return result


def extract_pptx(path: Path) -> dict:
    prs = Presentation(path)
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        item = {
            "slide": index,
            "texts": [],
            "tables": [],
            "charts": [],
            "images": 0,
            "notes": "",
        }
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                item["images"] += 1
            if getattr(shape, "has_text_frame", False):
                text = clean(shape.text_frame.text)
                if text:
                    item["texts"].append(text)
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append([clean(cell.text) for cell in row.cells])
                item["tables"].append(rows)
            if getattr(shape, "has_chart", False):
                item["charts"].append(extract_chart(shape.chart))
        try:
            notes = slide.notes_slide.notes_text_frame.text
            item["notes"] = clean(notes)
        except Exception:
            pass
        slides.append(item)
    return {"file": path.name, "kind": "pptx", "slide_count": len(slides), "slides": slides}


def extract_docx(path: Path) -> dict:
    doc = Document(path)
    blocks = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean(paragraph.text)
        if text:
            blocks.append({"type": "paragraph", "index": idx, "style": paragraph.style.name, "text": text})
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        blocks.append({"type": "table", "index": t_idx, "rows": rows})
    return {"file": path.name, "kind": "docx", "blocks": blocks}


def main() -> None:
    results = []
    for path in sorted(LECTURES.iterdir(), key=lambda p: p.name.casefold()):
        if path.suffix.lower() == ".pptx":
            results.append(extract_pptx(path))
        elif path.suffix.lower() == ".docx":
            results.append(extract_docx(path))
    OUT.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
